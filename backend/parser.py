import os
import fitz
import re
import pdfplumber
from docx import Document
from typing import Dict, List, Optional
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymorphy3 as pymorphy2

morph = pymorphy2.MorphAnalyzer()

RUSSIAN_CITIES = {
        'моск', 'санкт-петербур', 'новосибирск', 'екатеринбур', 'казан',
        'новгород', 'челябинск', 'самара', 'омск', 'ростов-на-дону',
        'уфа', 'красноярск', 'пермь', 'воронеж', 'волгоград', 'саратов', "сургут"
        # ... добавьте другие при необходимости
    }
class PDFDataExtractor:
    def __init__(self):
        self.parsed_data = {}
        self.all_documents = []
        self.frontend_data = {
            'договор': {
                'номер': '',
                'дата': ''
            },
            'адресат': {
                'фио': '',
                'должность': ''
            },
            'подписант': {
                'фио': '',
                'должность': ''
            },
            'экземпляры': 1 
        }

    def set_frontend_data(self, contract_number: str, contract_date: str, addressee: dict, signer: dict, copies: int = 1):
        """Устанавливает данные из фронтенда"""
        self.frontend_data['договор']['номер'] = contract_number
        self.frontend_data['договор']['дата'] = contract_date
        self.frontend_data['адресат'] = addressee
        self.frontend_data['подписант'] = signer
        self.frontend_data['экземпляры'] = copies  

    def extract_text_from_pdf(self, pdf_path: str) -> tuple:
        with fitz.open(pdf_path) as doc:
            first_page_text = doc[0].get_text("text")
            last_page_text = doc[-1].get_text("text") if len(doc) > 1 else ""
            num_pages = len(doc)
        return first_page_text, last_page_text, num_pages

    def format_title(self, title):
        words = title.split()
        formatted_words = []
        inside_quotes = False  
        quote_buffer = []  

        for i, word in enumerate(words):
            if word in {"ООО", "АО", "ЗАО", "ПАО", "ОАО"}:
                formatted_words.append(word)  
            elif i == 0:
                formatted_words.append(word.capitalize())  
            elif "«" in word: 
                inside_quotes = True
                quote_buffer = [word[1:].lower()]  
            elif "»" in word:  
                quote_buffer.append(word[:-1].lower())  
                formatted_words.append(f'«{" ".join(quote_buffer)}»')
                inside_quotes = False
            elif inside_quotes:
                quote_buffer.append(word.lower())
            else:
                formatted_words.append(word.lower()) 
        
        return " ".join(formatted_words)

    


    def extract_title_and_text(self, pdf_path: str) -> tuple:
        def capitalize_russian_city(text: str) -> str:
            """Заменяет корень города на правильный регистр (например: сургутского -> Сургутского)."""
            for city in RUSSIAN_CITIES:
                pattern = re.compile(city, flags=re.IGNORECASE)
                matches = list(pattern.finditer(text))
                for match in matches:
                    start, end = match.span()
                    matched_part = text[start:end]
                    text = text[:start] + city.title() + text[end:]
            return text

        def process_quoted_text(text: str) -> str:
            """Обрабатывает текст в кавычках, делая каждое слово с заглавной буквы"""
            if '«' in text and '»' in text:
                start = text.find('«') + 1
                end = text.find('»')
                quoted = text[start:end]
                # Разбиваем на слова и делаем каждое слово с заглавной буквы
                processed = ' '.join([word.capitalize() for word in quoted.split()])
                return text[:start] + processed + text[end:]
            return text

        with pdfplumber.open(pdf_path) as pdf:
            page = pdf.pages[0]  
            lines = page.extract_text().split("\n")[:-1] 
            lines = [
                line for line in lines
                if len(re.sub(r'[^А-Яа-яA-Za-z]', '', line)) > 2
            ]

            lines = [line.replace(".в ", "") for line in lines]
            special_abbrs = {"АСУ", "ТП", "ТР", "ГКС", "СЗСК", "КИИ", "ЗСК", "ЛПУМГ"}
            highlighted_lines = []
            is_title = False
            idx_ooo = 0

            for i, line in enumerate(lines):
                idx_ooo -= 1
                if line.isupper():
                    is_title = True
                
                # Обрабатываем строки с ООО
                if "ООО" in line:
                    idx_ooo += i
                    
                    # Обрабатываем предыдущую строку, если она содержит кавычки
                    if i > 0 and ('«' in lines[i-1] or '»' in lines[i-1]):
                        lines[i-1] = process_quoted_text(lines[i-1])
                    
                    # Разделяем строку на части до и после ООО
                    parts = line.split("ООО", 1)
                    if len(parts) == 2:
                        before_ooo = parts[0].strip()
                        after_ooo = parts[1].strip()
                        
                        # Обрабатываем текст в кавычках после ООО
                        after_ooo = process_quoted_text(after_ooo)
                        
                        text = before_ooo + ' ООО ' + after_ooo
                    else:
                        text = line
                    
                    highlighted_lines.append(text)
                    continue

                # Остальная обработка строк...
                if line.isupper() and sum(char.isdigit() for char in line) >= 3:
                    highlighted_lines.append(line)
                    continue
                if line.isupper():
                    highlighted_lines.append(line.lower())
                    continue
                if not line.isupper() and is_title:
                    highlighted_lines.append(line)

            # Обработка первой строки и строки с ООО
            if highlighted_lines:
                highlighted_lines[0] = highlighted_lines[0].capitalize()
            if idx_ooo + 1 < len(highlighted_lines):
                highlighted_lines[idx_ooo + 1] = highlighted_lines[idx_ooo + 1].capitalize()

            # Разделение на заголовок и текст
            title_lines = highlighted_lines[:idx_ooo + 1]
            text_lines = highlighted_lines[idx_ooo + 1:]

            # Дополнительная обработка аббревиатур и городов
            for i in range(len(title_lines)):
                line = title_lines[i]
                for abbr in special_abbrs:
                    pattern = re.compile(rf'\b{abbr.lower()}\b', re.IGNORECASE)
                    line = pattern.sub(abbr, line)
                line = capitalize_russian_city(line)
                title_lines[i] = line

            # Обработка оставшегося текста
            text_lines = [line.strip() for line in text_lines if line.strip()]
            
            return " ".join(title_lines), " ".join(text_lines)


    def extract_kt_data(self, first_page: str, last_page: str, pdf_path: str) -> dict:
        title, remaining_text = self.extract_title_and_text(pdf_path)
        
        match_uch = re.search(r'Уч\. №\s*(\S+)', last_page)
        match_sheets = re.search(r'Отпечатан 1 экз\. на (\d+) л\.', last_page)
        match_date = re.search(r'Дата:\s*(.+)', last_page)
        match_exemplars = re.search(r'экз\. №№\s*(\S+)', last_page)
        
        return {
            'названиедокумента': title,
            'текстдокумента': remaining_text,
            'учетныйномер': match_uch.group(1).strip() if match_uch else "",
            'количестволистов': match_sheets.group(1).strip() if match_sheets else "",
            'датадокумента': match_date.group(1).strip() if match_date else "",
            'экземпляры': match_exemplars.group(1).strip() if match_exemplars else "",
            'типдокумента': 'КТ'
        }

    def extract_sp_data(self, first_page: str, last_page: str, num_pages: int, pdf_path: str) -> dict:
        title, remaining_text = self.extract_title_and_text(pdf_path)
        
        match_inv = re.search(r'Инв\. № подл\.\s*(\S+) ', first_page)
        
        return {
            'названиедокумента': title,
            'текстдокумента': remaining_text,
            'инвентарныйномер': match_inv.group(1).strip() if match_inv else "",
            'количествостраниц': num_pages,
            'типдокумента': 'СП'
        }

    def parse_pdf(self, pdf_path: str) -> dict:
        first_page, last_page, num_pages = self.extract_text_from_pdf(pdf_path)
        is_commercial_secret = "Коммерческая тайна" in first_page
        
        if is_commercial_secret:
            data = self.extract_kt_data(first_page, last_page, pdf_path)
        else:
            data = self.extract_sp_data(first_page, last_page, num_pages, pdf_path)
        
        data.update({
            'номердоговора': self.frontend_data['договор']['номер'],
            'датадоговора': self.frontend_data['договор']['дата'],
            'адресат': self.frontend_data['адресат'],
            'подписант': self.frontend_data['подписант']
        })

        self.all_documents.append(data)
        return data

    def generate_documents_list(self):
        ks_list = []
        # sp_list = []

         
        fixed_item_added = False 
        for doc in self.all_documents:
            if not fixed_item_added:
                fixed_item = f"Электронная версия {{documentation_type_second}} документации по проекту «{doc['названиедокумента']} на электронном носителе CD-R, коммерческая тайна, уч. № КТ/Э 459 от 26.03.2025(календарь), экз. № 1/1(человек должен задать отдельно для дисков и бумажной документации) только в адрес. "
                ks_list.append(fixed_item)
                fixed_item_added = True

            if doc['типдокумента'] == 'КТ':
                ks_entry = f"{doc['текстдокумента']}, коммерческая тайна, Уч. № {doc['учетныйномер']} от {doc['датадокумента']}, экз. №№ 1/1-1/{self.frontend_data['экземпляры']} на {doc['количестволистов']} л. каждый, только в адрес."
                ks_list.append(ks_entry)
            elif doc['типдокумента'] == 'СП':
                sp_entry = f"{doc['текстдокумента']}, инв. № {doc['инвентарныйномер']} - на {doc['количествостраниц']} л. в  {self.frontend_data['экземпляры']} экз."
                ks_list.append(sp_entry)   #sp_list.append(sp_entry)
        
        # result = ""
        # if ks_list:
        #     result += "\r\n\r\n".join(ks_list)
        # if sp_list:
        #     result += "\r\n".join(sp_list)
        
        return ks_list #result

class DocxTemplateProcessor:
    def __init__(self):
        self.data_sources = {
            'парсера': {},
            'бд': {},
            'списокдокументов': "",
            'формы':{}
        }

    def add_frontend_data(self, data: dict):
        """Добавляет данные из фронтенда"""
        self.data_sources['формы'] = data

    def add_parser_data(self, data: dict):
        """Добавляет данные из парсера."""
        self.data_sources['парсера'] = data

    def add_db_data(self, data: dict):
        """Добавляет данные из базы данных."""
        self.data_sources['бд'] = data

    def add_documents_list(self, documents_list: str):
        """Добавляет список документов."""
        self.data_sources['списокдокументов'] = documents_list

    def _format_fio(self, fio: str, format_key: str) -> str:
        """Форматирует ФИО в зависимости от регистра символов в ключе"""
        if not fio:
            return ""
        
        parts = [p.strip() for p in fio.split() if p.strip()]
        if not parts:
            return ""
        
        for i in range(1, len(parts)):
            parsed = morph.parse(parts[i])[0]
            nominative = parsed.inflect({'nomn'})  # именительный падеж
            if nominative:
                parts[i] = nominative.word.capitalize()

        # Получаем часть ключа до "_из_формы" и убираем фигурные скобки
        format_part = format_key.split('_')[0].split('.')[-1].replace('{', '').replace('}', '')

        match format_part:
            case 'ФИО':
                return " ".join(parts)  # Полное ФИО: Фамилия Имя Отчество
            case 'ФИо':
                if len(parts) >= 3:
                    return f"{parts[0]} {parts[1]} {parts[2][0]}."
                elif len(parts) == 2:
                    return f"{parts[0]} {parts[1]}"
                return parts[0]
            case 'Фио':
                if len(parts) >= 3:
                    return f"{parts[0]} {parts[1][0]}.{parts[2][0]}."
                elif len(parts) == 2:
                    return f"{parts[0]} {parts[1][0]}."
                return parts[0]
            case 'иоФ':
                if len(parts) >= 3:
                    return f"{parts[1][0]}.{parts[2][0]}. {parts[0]}"
                elif len(parts) == 2:
                    return f"{parts[1][0]}. {parts[0]}"
                return parts[0]
            case 'ИО':
                if len(parts) >= 3:
                    return f"{parts[1]} {parts[2]}"
                elif len(parts) >= 2:
                    return parts[1]
                return parts[0]
            case _:
                return " ".join(parts)  # По умолчанию — полное ФИО
        
    def process_template(self, template_path: str, output_path: str):
        """Обрабатывает шаблон с выравниванием по ширине и заданными отступами"""
        doc = Document(template_path)
        
        # Параметры в twips (1 см = 567 twips)
        LEFT_INDENT = int(0.25 * 567)    # 0.25 см
        FIRST_LINE_INDENT = int(1.0 * 567)  # 1.0 см
        TAB_STOP_1 = int(1.75 * 567)      # 1.75 см (первая позиция табуляции)
        LINE_SPACING = 360                # 1.5 строки
        
        # Ищем параграф с меткой {списокдокументов:}
        for paragraph in doc.paragraphs:
            if '{списокдокументов:}' in paragraph.text:
                parent = paragraph._p.getparent()
                index = parent.index(paragraph._p)
                parent.remove(paragraph._p)
                
                for item in self.data_sources['списокдокументов']:
                    new_paragraph = OxmlElement('w:p')
                    
                    # Настройки нумерации
                    num_pr = OxmlElement('w:numPr')
                    ilvl = OxmlElement('w:ilvl')
                    ilvl.set(qn('w:val'), '0')
                    num_id = OxmlElement('w:numId')
                    num_id.set(qn('w:val'), '1')
                    num_pr.append(ilvl)
                    num_pr.append(num_id)
                    
                    # Настройки абзаца
                    p_pr = OxmlElement('w:pPr')
                    p_pr.append(num_pr)
                    
                    # Выравнивание по ширине
                    jc = OxmlElement('w:jc')
                    jc.set(qn('w:val'), 'both')
                    p_pr.append(jc)
                    
                    # Отступы (левый 0.25 см, первая строка 1 см)
                    ind = OxmlElement('w:ind')
                    ind.set(qn('w:left'), str(LEFT_INDENT))
                    ind.set(qn('w:firstLine'), str(FIRST_LINE_INDENT))
                    p_pr.append(ind)
                    
                    # Позиции табуляции (1.75 см)
                    tabs = OxmlElement('w:tabs')
                    tab = OxmlElement('w:tab')
                    tab.set(qn('w:val'), 'left')
                    tab.set(qn('w:pos'), str(TAB_STOP_1))
                    tabs.append(tab)
                    p_pr.append(tabs)
                    
                    # Настройки интервала
                    spacing = OxmlElement('w:spacing')
                    spacing.set(qn('w:before'), '0')
                    spacing.set(qn('w:after'), '0')
                    spacing.set(qn('w:line'), str(LINE_SPACING))
                    spacing.set(qn('w:lineRule'), 'auto')
                    p_pr.append(spacing)
                    
                    new_paragraph.append(p_pr)
                    
                    # Добавляем текст
                    run = OxmlElement('w:r')
                    text = OxmlElement('w:t')
                    text.text = item
                    run.append(text)
                    new_paragraph.append(run)
                    
                    parent.insert(index, new_paragraph)
                    index += 1
                    
                break
        
        # Остальная обработка документа
        self._process_headers_footers(doc)

        for paragraph in doc.paragraphs:
            self._replace_in_paragraph(paragraph)
            
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph)

        for shape in doc.inline_shapes:
            if hasattr(shape, 'text_frame'):
                for paragraph in shape.text_frame.paragraphs:
                    self._replace_in_paragraph(paragraph)

        doc.save(output_path)  
   
        

    def _process_headers_footers(self, doc):
        """Обрабатывает верхние и нижние колонтитулы всех секций."""
        for section in doc.sections:
            for header_part in [section.header, section.footer, 
                              section.even_page_header, section.even_page_footer,
                              section.first_page_header, section.first_page_footer]:
                if header_part is not None:
                    for paragraph in header_part.paragraphs:
                        self._replace_in_paragraph(paragraph)
                    for table in header_part.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                for paragraph in cell.paragraphs:
                                    self._replace_in_paragraph(paragraph)

    def _replace_in_paragraph(self, paragraph):
        """Заменяет ключи с учетом любых форматов ФИО"""
        keys = self._find_all_keys(paragraph.text)
        if not keys:
            return

        runs = paragraph.runs
        full_text = ''.join(run.text for run in runs)
        
        for key, key_info in keys.items():
            if not key_info:
                continue
                
            replacement = ""
            source_data = self.data_sources.get(key_info['source'], {})
            
            # Обработка всех вариантов ключей ФИО
            if any(fio_key in key.lower() for fio_key in ['фио', 'иоф', 'ио', 'фИО', 'ФИО']):
                # Для ключей вида {ФИО_из_бд}
                if key_info['source'] == 'бд' and key_info.get('object', '').upper() == 'ФИО':
                    replacement = str(source_data.get(key_info['object'], ""))
                
                # Для вложенных ключей {адресат.Фио_из_формы}
                elif 'object' in key_info:
                    obj = source_data.get(key_info['object'], {})
                    if isinstance(obj, dict):
                        fio_value = str(obj.get('фио', ""))
                    else:
                        fio_value = str(obj) if obj else ""
                    replacement = self._format_fio(fio_value, key)
            
            # Остальные ключи
            elif 'object' in key_info:
                if key_info.get('field'):
                    obj = source_data.get(key_info['object'], {})
                    if isinstance(obj, dict):
                        replacement = str(obj.get(key_info['field'], ""))
                    else:
                        replacement = str(obj) if obj else ""
                else:
                    replacement = str(source_data.get(key_info['object'], ""))
            elif '.' in key_info.get('key', ''):
                parts = key_info['key'].split('.')
                current = source_data
                for part in parts:
                    if isinstance(current, dict) and part in current:
                        current = current[part]
                    else:
                        current = ""
                        break
                replacement = str(current)
            else:
                replacement = str(source_data.get(key_info.get('key', ""), ""))

            print(f"Processing key: {key} → {replacement}")
            full_text = full_text.replace(key, replacement)

        # Обновляем текст
        for run in runs:
            run.text = ""
        if runs:
            runs[0].text = full_text
    
    def _copy_run_formatting(self, source_run, target_run):
        """Копирует форматирование из одного run в другой."""
        target_run.bold = source_run.bold
        target_run.italic = source_run.italic
        target_run.underline = source_run.underline
        target_run.font.name = source_run.font.name
        target_run.font.size = source_run.font.size
        if source_run.font.color.rgb:
            target_run.font.color.rgb = source_run.font.color.rgb

    def _find_all_keys(self, text: str) -> Dict[str, dict]:
        keys = {}
        
        # 1. Обрабатываем сложные ключи вида {адресат.Фио_из_формы}
        complex_pattern = re.compile(r'\{([^}.]+)(?:\.([^}_]+))?_из_([^}]+)\}')
        for match in complex_pattern.finditer(text):
            full_key = match.group(0)
            object_name = match.group(1)  # "адресат"
            field_name = match.group(2)   # "Фио"
            source = match.group(3)       # "формы"
            
            keys[full_key] = {
                'found': True,
                'source': source,
                'object': object_name,
                'field': field_name.lower() if field_name else None # Приводим к нижнему регистру для унификации
            }

        # 2. Обрабатываем простые ключи вида {номердоговора}
        simple_pattern = re.compile(r'\{([^}:]+)(?::)?\}')
        for match in simple_pattern.finditer(text):
            full_key = match.group(0)
            key_name = match.group(1)
            
            # Проверяем плоские ключи во всех источниках
            for source_name, source_data in self.data_sources.items():
                if key_name in source_data and not isinstance(source_data[key_name], (dict, list)):
                    keys[full_key] = {
                        'found': True,
                        'source': source_name,
                        'key': key_name
                    }
                    break
                
                # Проверяем вложенные структуры (например, данные договора)
                if isinstance(source_data, dict) and '.' in key_name:
                    parts = key_name.split('.')
                    current = source_data
                    valid = True
                    for part in parts:
                        if part in current:
                            current = current[part]
                        else:
                            valid = False
                            break
                    if valid:
                        keys[full_key] = {
                            'found': True,
                            'source': source_name,
                            'key': key_name
                        }
                        break
        
        return keys
 

def main(
        docx_path: str,
        recipient: str,
        signer: str,
        contract_number: Optional[str] = None,
        contract_date: Optional[str] = None,
        pdf_folder_path: Optional[str] = None
        ) -> str:
    pdf_extractor = PDFDataExtractor()
    docx_processor = DocxTemplateProcessor()

    addressee = {
        'фио': recipient,
        'должность': 'Заместителю генерального директора',
        'организация': 'ООО «Мяу Кусь»'
    }
    signer = {
        'фио': signer,
        'должность': 'Генеральный директор'
    }
    copies = 4
    
    pdf_extractor.set_frontend_data(
        contract_number,
        contract_date,
        addressee,
        signer,
        copies
    )

    docx_processor.add_frontend_data({
        'договор': {
            'номер': contract_number,
            'дата': contract_date
        },
        'адресат': addressee,
        'подписант': signer,
        'экземпляры': copies 
    })

    pdf_files = []
    pdf_folder_path = r"/app/pdfs"
    if pdf_folder_path:
        if not os.path.exists(pdf_folder_path):
            raise ValueError(f"Путь к PDF не существует: {pdf_folder_path}")
            
        for root, _, files in os.walk(pdf_folder_path):
            pdf_files.extend([os.path.join(root, f) for f in files if f.lower().endswith('.pdf')])
        
        print(f"Найдено {len(pdf_files)} PDF файлов. Начинаю парсинг...")
        
        for pdf_file in pdf_files:
            print(f"Парсинг файла: {pdf_file}")
            parsed_data = pdf_extractor.parse_pdf(pdf_file)
            docx_processor.add_parser_data(parsed_data)
        
        documents_list = pdf_extractor.generate_documents_list()
        docx_processor.add_documents_list(documents_list)
        print("\nСформированный список документов:\n")
        print(documents_list) 

    template_path = f"static{docx_path}"
    print(template_path)
    
    original_filename = os.path.basename(docx_path)
    output_path = os.path.join("static/converted_files", original_filename)
    docx_processor.process_template(template_path, output_path)
    print(f"\nДокумент сохранен как {output_path}")
    return output_path

if __name__ == "__main__":
    main()