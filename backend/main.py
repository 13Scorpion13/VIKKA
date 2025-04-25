from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import List
from bs4 import BeautifulSoup
from docx import Document
import os
import re
from datetime import datetime
import time
from pathlib import Path
from parser import main

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static/preview", StaticFiles(directory="./static/preview"), name="preview")
app.mount("/static/converted_files", StaticFiles(directory="./static/converted_files"), name="converted_files")
app.mount("/static/templates", StaticFiles(directory="./static/templates"), name="templates")

class Template(BaseModel):
    id: int
    title: str
    preview_image: str
    document_path: str
    last_modified: str = None

templates_db = [
    {"id": 1, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 2, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template1.docx", "last_modified": "5 минут назад"},
    {"id": 3, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 4, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 5, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 6, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 7, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 8, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 9, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 10, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 11, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 12, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 13, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
    {"id": 14, "title": "О представлении ", "preview_image": "preview/image2.png", "document_path": "templates/template.docx", "last_modified": "5 минут назад"},
]

@app.get("/api/templates/recent", response_model=List[Template])
async def get_recent_templates(limit: int = 100):
    return templates_db[:limit]

@app.get("/api/templates/all", response_model=List[Template])
async def get_all_templates(skip: int = 0, limit: int = 100):
    return templates_db[skip : skip + limit]

@app.get("/api/templates/search")
async def search_templates(query: str):
    return [t for t in templates_db if query.lower() in t["title"].lower()]

@app.get("/images/{image_path:path}")
async def get_image(image_path: str):
    image_full_path = os.path.join("static", image_path)
    print(image_full_path)
    if not os.path.exists(image_full_path):
        raise HTTPException(status_code=404, detail="Image not found")
    return FileResponse(image_full_path)

@app.post("/api/process-document/")
async def process_document(request: Request):
    try:
        data = await request.json()
        print("Полученные данные:", data)

        
        contract_number = data.get("contract_number")
        contract_date = data.get("contract_date")
        recipient = data.get("recipient")
        signer = data.get("signer")
        pdf_folder_path = data.get("pdf_folder_path")
        docx_path = data.get("docx_path")

        if not docx_path:
            raise HTTPException(status_code=400, detail="Не указан путь к документу")
        
        print(f"[1/6] Получен путь к DOCX: {docx_path}")

        full_docx_path = main(
            docx_path,
            contract_number,
            contract_date,
            recipient,
            signer,
            pdf_folder_path
        )
        print(f"[2/6] Обработанный путь: {full_docx_path}")
        
        
        return JSONResponse({
            "full_docx_path": full_docx_path,
            "status": "success"
        })
        
    except HTTPException:
        raise
    except Exception as e:
        error_msg = f"Непредвиденная ошибка: {str(e)}"
        print(f"[CRITICAL ERROR] {error_msg}")
        raise HTTPException(status_code=500, detail=error_msg)
    
@app.post("/api/extract-fields")
async def extract_fields(request: Request):
    """Извлекает поля вида {field_name} из DOCX."""
    try:
        data = await request.json()
        docx_path = data.get("docx_path")
        
        if not docx_path:
            raise HTTPException(status_code=400, detail="Не указан путь к документу")
        
        doc = Document(docx_path)
        fields = set()
        
        
        for paragraph in doc.paragraphs:
            matches = re.findall(r"\{(\w+)\}", paragraph.text)
            fields.update(matches)
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    matches = re.findall(r"\{(\w+)\}", cell.text)
                    fields.update(matches)
        
        return {"fields": list(fields)}
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка при извлечении полей: {str(e)}")
    
@app.post("/api/replace-fields")
async def replace_fields(request: Request):
    try:
        data = await request.json()
        docx_path = data.get("docx_path")
        fields_data = data.get("fields_data", {})
        
        if not docx_path:
            raise HTTPException(status_code=400, detail="Не указан путь к документу")
        
        if not os.path.exists(docx_path):
            raise HTTPException(status_code=404, detail="Файл не найден")
        
        doc = Document(docx_path)
        

        for paragraph in doc.paragraphs:
            for field, value in fields_data.items():
                if f"{{{field}}}" in paragraph.text:
                    paragraph.text = paragraph.text.replace(f"{{{field}}}", str(value))
        
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for field, value in fields_data.items():
                        if f"{{{field}}}" in cell.text:
                            cell.text = cell.text.replace(f"{{{field}}}", str(value))
        
        
        
        doc.save(docx_path)
        
        return {"status": "success", "message": "Файл успешно обновлен", "file_path": docx_path}
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка при замене полей: {str(e)}")