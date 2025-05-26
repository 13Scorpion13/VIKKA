from datetime import datetime
import os
import re
import httpx
import shutil
from typing import List
from docx import Document
from fastapi import Depends, HTTPException, APIRouter, Request, UploadFile, File, Form
from fastapi.responses import JSONResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import insert, select, update, delete, desc, func
from parser import main
from database import get_async_session
from models import template, letter_history
from schemas import TemplateCreate, TemplateUpdate, TemplateRead
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from pathlib import Path
from db_utils import get_adressee_by_id, get_signer_by_id, get_employee_by_id
# from adressee_router import get_adressee_by_id
# from signer_router import get_signer_by_id

UPLOAD_FOLDER = "/app/uploads"
router = APIRouter(prefix="/api", tags=["api"])

templates_db = [
    {
        "id": 1,
        "title": "О представлении документации",
        "preview_image": "/image2.png",
        "document_path": "/templates/Письмо о представлении документации.docx",
        "form_type": "full",
        "last_modified": "5 минут назад"
    },
    {
        "id": 2,
        "title": "Письмо о допуске ",
        "preview_image": "/image2.png",
        "document_path": "/templates/Письмо о допуске.docx",
        "form_type": "minimal",
        "last_modified": "5 минут назад"
    }
]

""" @router.get("/templates/recent", response_model=List[TemplateRead])
async def get_recent_templates(limit: int = 100, session: AsyncSession = Depends(get_async_session)):
    query = select(template).limit(limit)
    result = await session.execute(query)
    
    # Получаем список словарей
    rows = result.mappings().all()

    transformed = []
    transformed = [
        dict(row) | {"last_modified": row["last_modified"].strftime("%Y-%m-%d")}
        for row in rows
    ]

    # transformed = [
    #     {
    #         "id": row.id,
    #         "title": row.title,
    #         "preview_image": row.preview_image,
    #         "document_path": row.document_path,
    #         "last_modified": humanize.naturaltime(datetime.now() - row.last_modified),
    #         "form_type": row.form_type
    #     }
    #     for row in rows
    # ]

    return transformed[: limit] """
@router.get("/templates/recent/{user_id}")
async def get_recent_templates(
    user_id: int,
    limit: int = 100,
    session: AsyncSession = Depends(get_async_session)
):
    # Подзапрос: последние даты использования шаблонов для данного пользователя
    latest_dates_subquery = (
        select(
            letter_history.c.template_id,
            func.max(letter_history.c.date).label("latest_date")
        )
        .where(letter_history.c.user_id == user_id)
        .group_by(letter_history.c.template_id)
        .subquery()
    )

    # Основной запрос: выбираем шаблоны + дату последнего использования
    query = (
        select(template, latest_dates_subquery.c.latest_date)
        .join(latest_dates_subquery, template.c.id == latest_dates_subquery.c.template_id)
        .limit(limit)
    )

    result = await session.execute(query)
    rows = result.mappings().all()

    # Форматируем ответ
    transformed = [
        dict(row) | {
            "last_modified": row["latest_date"].strftime("%Y-%m-%d")
        }
        for row in rows
    ]

    return transformed[:limit]

@router.get("/templates/all", response_model=List[TemplateRead])
async def get_all_templates(
    skip: int = 0,
    limit: int = 100,
    session: AsyncSession = Depends(get_async_session)
):
    query = select(template).limit(limit)
    result = await session.execute(query)
    
    # Получаем список словарей
    rows = result.mappings().all()

    # Преобразуем last_modified в нужный формат
    transformed = [
        dict(row) | {"last_modified": row["last_modified"].strftime("%Y-%m-%d")}
        for row in rows
    ]
    return transformed

    # return templates_db[skip : skip + limit]

@router.get("/templates/search")
async def search_templates(query: str):
    return [t for t in templates_db if query.lower() in t["title"].lower()]

@router.post("/process-document/")
async def process_document(
    request: Request,
    session: AsyncSession = Depends(get_async_session)
):
    try:
        data = await request.json()
        # pdf_folder_path = "/app/uploads"
        print("Полученные данные:", data)

        if not all([data.get("docx_path"), data.get("adressee_id"), data.get("signer_id"), data.get("user_id")]):
            raise HTTPException(status_code=400, detail="Обязательные поля отсутствуют")
        
        db_adressee = await get_adressee_by_id(session, data["adressee_id"])
        print(db_adressee)
        db_signer = await get_signer_by_id(session, data["signer_id"])
        print(db_adressee)

        if not db_adressee or not db_signer:
            raise HTTPException(status_code=404, detail="Адресат или подписант не найдены")
        
        pdf_folder_path = f"/app/uploads/{data['user_id']}"

        # Формируем данные для main()
        processed_data = {
            "docx_path": data["docx_path"],
            "user_id": data["user_id"],
            # "recipient": db_adressee["full_name"],
            # "signer": db_signer["full_name"],
            "adressee_data": {
                "фио": db_adressee["full_name"],
                "должность": db_adressee["position"],
                "организация": db_adressee["organization"]
            },
            "signer_data": {
                "фио": db_signer["full_name"],
                "должность": db_signer["position"]
            }
        }

        # recipient = data.get("recipient")
        # signer = data.get("signer")
        # docx_path = data.get("docx_path")

        # if not all([docx_path, recipient, signer]):
        #     raise HTTPException(status_code=400, detail="Обязательные поля отсутствуют")
        
        # template = next((t for t in templates_db if t["document_path"] == docx_path), None)
        # form_type = template.get("form_type") if template else "full"

        form_type = next((t["form_type"] for t in templates_db if t["document_path"] == data["docx_path"]), "full")

        if form_type == "full":
            processed_data.update({
                "contract_number": data.get("contract_number"),
                "contract_date": data.get("contract_date"),
                "pdf_folder_path": pdf_folder_path
            })
        elif form_type == "minimal":
            processed_data.update({
                "organisation": data.get("organisation"),
                "contract": data.get("contract")
            })

        full_docx_path = main(**processed_data)

        # if form_type == "full":
        #     contract_number = data.get("contract_number")
        #     contract_date = data.get("contract_date")
        #     pdf_folder_path = data.get("pdf_folder_path")
            
        #     if not all([contract_number, contract_date, pdf_folder_path]):
        #         raise HTTPException(status_code=400, detail="Не заполнены обязательные поля для полной формы")
            
        # if form_type == "minimal":
        #     organisation = data.get("organisation")
        #     contract = data.get("contract")

        #     if not all([organisation, contract]):
        #         raise HTTPException(status_code=400, detail="Не заполнены обязательные поля для полной формы")
        
        # full_docx_path = main(
        #     docx_path=data["docx_path"],
        #     recipient=data["recipient"],
        #     signer=data["signer"],
        #     contract_number=data.get("contract_number"),
        #     contract_date=data.get("contract_date"),
        #     pdf_folder_path=data.get("pdf_folder_path"),
        #     organisation = data.get("organisation"),
        #     contract = data.get("contract")
        # )
        print(f"[2/6] Обработанный путь: {full_docx_path}")
        
        
        return JSONResponse({
            "full_docx_path": full_docx_path,
            "status": "success",
            "recipient": db_adressee["full_name"],
            "signer": db_signer["full_name"]
        })
        
    except HTTPException:
        raise
    except Exception as e:
        error_msg = f"Непредвиденная ошибка: {str(e)}"
        print(f"[CRITICAL ERROR] {error_msg}")
        raise HTTPException(status_code=500, detail=error_msg)
    
@router.post("/extract-fields")
async def extract_fields(request: Request):
    """Извлекает поля вида {field_name} из DOCX."""
    try:
        data = await request.json()
        docx_path = data.get("docx_path")
        
        if not docx_path:
            raise HTTPException(status_code=400, detail="Не указан путь к документу")
        
        doc = Document(docx_path)
        fields = set()

        def extract_fields_from_text(text):
            # Находим все поля вида {field_name} или {field_name_N}
            return re.findall(r"\{(\w+)(?:_\d+)?\}", text)
        
        for paragraph in doc.paragraphs:
            fields.update(extract_fields_from_text(paragraph.text))
    
        # Обрабатываем headers и footers
        for section in doc.sections:
            for paragraph in section.header.paragraphs:
                fields.update(extract_fields_from_text(paragraph.text))
            for paragraph in section.footer.paragraphs:
                fields.update(extract_fields_from_text(paragraph.text))
        
        # Обрабатываем таблицы (здесь нам нужны полные имена с индексами)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    # Для таблиц используем полный вариант с индексами
                    matches = re.findall(r"\{(\w+_\d+)\}", cell.text)
                    if not matches:
                        # Если нет индексированных, проверяем обычные поля
                        matches = re.findall(r"\{(\w+)\}", cell.text)
                    fields.update(matches)

        print("Извлеченные поля:", fields)
        
        
        # for paragraph in doc.paragraphs:
        #     fields.update(re.findall(r"\{(\w+)\}", paragraph.text))
    
        # for section in doc.sections:
        #     for paragraph in section.header.paragraphs:
        #         fields.update(re.findall(r"\{(\w+)\}", paragraph.text))
        #     for paragraph in section.footer.paragraphs:
        #         fields.update(re.findall(r"\{(\w+)\}", paragraph.text))
        
        # for table in doc.tables:
        #     for row in table.rows:
        #         for cell in row.cells:
        #             matches = re.findall(r"\{(\w+)\}", cell.text)
        #             fields.update(matches)
        
        return {"fields": list(fields)}
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка при извлечении полей: {str(e)}")
    
@router.post("/replace-fields")
async def replace_fields(request: Request):
    try:
        data = await request.json()
        docx_path = data.get("docx_path")
        print(f"Путь к файлу для сохранения: {docx_path}")
        fields_data = data.get("fields_data", {})
        contract_field = fields_data.get("contract_field")
        data_field = fields_data.get("data_field")
        user_id = data.get("user_id")

        
        if not docx_path:
            raise HTTPException(status_code=400, detail="Не указан путь к документу")
        
        if not os.path.exists(docx_path):
            raise HTTPException(status_code=404, detail="Файл не найден")
        
        doc = Document(docx_path)

        for paragraph in doc.paragraphs:
            for field, value in fields_data.items():
                if f"{{{field}}}" in paragraph.text:
                    paragraph.text = paragraph.text.replace(f"{{{field}}}", str(value))
        
        
        for table in doc.tables:
            for row in table.rows:
                count_cell = 0
                for cell in row.cells:
                    count_cell += 1
                    for field, value in fields_data.items():
                        print(f"field: {field}, value: {value}")
                        if f"{{{field}}}" in cell.text:
                            cell.text = cell.text.replace(f"{{{field}}}", str(value))
                            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            for paragraph in cell.paragraphs:
                                if len(table.rows[0].cells) == 4:
                                    if count_cell == 4:
                                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                    else:
                                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                else:
                                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                    for run in paragraph.runs:
                                        run.font.italic = True

        docx_path = f"static/converted_files/{user_id}/Исх. № {contract_field} от {data_field} {os.path.basename(docx_path)}"
        print(docx_path)
        doc.save(docx_path)
        # print(int(data.get("user_id")), int(data.get("templateID")), docx_path, datetime.now().isoformat(), sep='\n')
        try:
            async with httpx.AsyncClient() as client:
                await client.post(
                    "http://localhost:8000/letter_history/",
                    json={
                        "user_id": data.get("user_id"), 
                        "signer_id": 1,
                        "template_id": int(data.get("templateID")),
                        "file_path": docx_path,
                        "adressee_id": 1
                    }
                )
        except Exception as e:
            print(f"ERROR!!! {e}")
        return {"status": "success", "message": "Файл успешно обновлен", "file_path": docx_path}
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка при замене полей: {str(e)}")
    
# @router.post("/add-table-row")
# async def add_table_row(request: Request):
#     data = await request.json()
#     doc = Document(data["docx_path"])
#     table = doc.tables[data["table_index"]]

#     row_index = len(table.rows) - 1
#     print(f"row_index: {row_index}")
    
#     #row_number = len(table.rows)
#     new_row = table.add_row()
#     print(f"new_row: {new_row}")
    
#     if len(new_row.cells) >= 4:
#         new_row.cells[0].text = str(row_index + 1)
#         new_row.cells[1].text = f"{{engineer_full_name_{row_index}}}, {{engineer_position_{row_index}}}"
#         new_row.cells[2].text = f"{{engineer_notebook_{row_index}}}"
#         new_row.cells[3].text = f"{{areas_name}}"
        
#         for cell in new_row.cells:
#             for paragraph in cell.paragraphs:
#                 paragraph.paragraph_format.space_after = Pt(0)
#                 paragraph.paragraph_format.line_spacing = Pt(40)
#                 paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
#     doc.save(data["docx_path"])
#     return {"status": "success", "file_path": data["docx_path"]}

def merge_last_column(table):
    """
    Объединяет все ячейки в последнем столбце таблицы.
    """
    if len(table.rows) < 1:
        return

    col_index = len(table.columns) - 1  # индекс последнего столбца
    first_cell = table.rows[1].cells[col_index]
    last_cell = table.rows[-1].cells[col_index]

    # Проверяем, не объединены ли они уже
    if first_cell._tc is not last_cell._tc:
        merged_cell = first_cell.merge(last_cell)
        merged_cell.text = "{areas_name}"  # устанавливаем макрос
        merged_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        merged_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

@router.post("/add-table-row")
async def add_table_row(request: Request):
    data = await request.json()
    doc = Document(data["docx_path"])
    table = doc.tables[data["table_index"]]

    row_index = len(table.rows) - 1
    new_row = table.add_row()

    if len(new_row.cells) >= 4:
        new_row.cells[0].text = str(row_index + 1)
        new_row.cells[1].text = f"{{engineer_full_name_{row_index}}}, \n{{engineer_position_{row_index}}}"
        new_row.cells[2].text = f"{{engineer_notebook_{row_index}}}"

        for cell in new_row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(40)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Перед сохранением — объединяем последний столбец
    merge_last_column(table)

    doc.save(data["docx_path"])
    return {"status": "success", "file_path": data["docx_path"]}

@router.get("/engineers")
async def get_engineers():
    engineers = [
        {"full_name": "Петров Пётр Петрович", "notebook": "Dell XPS, \nS/N:467840042", "position": "Ассистент"},
        {"full_name": "Иванов Иван Иванович", "notebook": "Lenovo X1, \nS/N:467840042", "position": "Ведущий инженер"},
        {"full_name": "Сидоров Савелий Савельевич", "notebook": "MacBook Pro, \nS/N:467840042", "position": "Инженер"},
        {"full_name": "Иванов Иван Иванович", "notebook": "Lenovo X1, \nS/N:467840042", "position": "Ведущий инженер"},
    ]
    return {"engineers": engineers}

@router.get("/areas")
async def get_areas():
    areas = [
        {"name": "Зона A"}, 
        {"name": "Зона B"},
        {"name": "Зона C"},
        {"name": "Зона D"},
    ]
    return {"areas": areas}

@router.post("/upload-files/")
async def upload_files(
    files: list[UploadFile] = File(...),
    user_id: int = Form(...)
):
    try:
        user_folder = Path(UPLOAD_FOLDER) / str(user_id)
        user_folder.mkdir(parents=True, exist_ok=True)
        # Path(UPLOAD_FOLDER).mkdir(parents=True, exist_ok=True)

        for file in user_folder.glob("*"):
            try:
                if file.is_file():
                    file.unlink()
            except Exception as e:
                print(f"Error deleting {file}: {e}")
        
        # for filename in os.listdir(UPLOAD_FOLDER):
        #     file_path = os.path.join(UPLOAD_FOLDER, filename)
        #     try:
        #         if os.path.isfile(file_path):
        #             os.unlink(file_path)
        #     except Exception as e:
        #         print(f"Ошибка при удалении файла {file_path}: {e}")

        # Сохраняем файлы
        # saved_files = []
        # for file in files:
        #     try:
        #         file_path = os.path.join(UPLOAD_FOLDER, file.filename)
        #         with open(file_path, "wb") as buffer:
        #             content = await file.read()
        #             buffer.write(content)
        #         saved_files.append(file.filename)
        #     except Exception as e:
        #         print(f"Ошибка при сохранении файла {file.filename}: {e}")
        #         continue

        saved_files = []
        for file in files:
            try:
                file_path = user_folder / file.filename
                with file_path.open("wb") as buffer:
                    content = await file.read()
                    buffer.write(content)
                saved_files.append(file.filename)
            except Exception as e:
                print(f"Error saving {file.filename}: {e}")
                continue

        # return {
        #     "status": "success",
        #     "files": saved_files,
        #     "path": UPLOAD_FOLDER  # Возвращаем путь внутри контейнера
        # }
        return {
                "status": "success",
                "files": saved_files,
                "user_id": user_id
            }
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Ошибка загрузки файлов: {str(e)}"
        )