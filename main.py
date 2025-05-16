from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import List
from bs4 import BeautifulSoup
from docx import Document
import os
import re
from datetime import datetime
import time
from pathlib import Path
from parser import main
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/preview", StaticFiles(directory="preview"), name="preview")
app.mount("/converted_files", StaticFiles(directory="converted_files"), name="converted_files")

class Template(BaseModel):
    id: int
    title: str
    preview_image: str
    document_path: str
    form_type: str
    last_modified: str = None

templates_db = [
    {
        "id": 1,
        "title": "О представлении документации",
        "preview_image": "/image2.png",
        "document_path": "/templates/Письмо о представлении документации.docx",
        "form_type": "full",
        "last_modified": "5 минут назад"
    },
    {
        "id": 2,
        "title": "Письмо о допуске ",
        "preview_image": "/image2.png",
        "document_path": "/templates/Письмо о допуске.docx",
        "form_type": "minimal",
        "last_modified": "5 минут назад"
    }
]

@app.get("/api/templates/recent", response_model=List[Template])
async def get_recent_templates(limit: int = 100):
    return templates_db[:limit]

@app.get("/api/templates/all", response_model=List[Template])
async def get_all_templates(skip: int = 0, limit: int = 100):
    return templates_db[skip : skip + limit]

@app.get("/api/templates/search")
async def search_templates(query: str):
    return [t for t in templates_db if query.lower() in t["title"].lower()]

@app.get("/images/{image_path:path}")
async def get_image(image_path: str):
    image_full_path = os.path.join("static", image_path)
    if not os.path.exists(image_full_path):
        raise HTTPException(status_code=404, detail="Image not found")
    return FileResponse(image_full_path)

@app.post("/api/process-document/")
async def process_document(request: Request):
    try:
        data = await request.json()
        print("Полученные данные:", data)

        recipient = data.get("recipient")
        signer = data.get("signer")
        docx_path = data.get("docx_path")

        if not all([docx_path, recipient, signer]):
            raise HTTPException(status_code=400, detail="Обязательные поля отсутствуют")
        
        template = next((t for t in templates_db if t["document_path"] == docx_path), None)
        form_type = template.get("form_type") if template else "full"

        if form_type == "full":
            contract_number = data.get("contract_number")
            contract_date = data.get("contract_date")
            pdf_folder_path = data.get("pdf_folder_path")
            
            if not all([contract_number, contract_date, pdf_folder_path]):
                raise HTTPException(status_code=400, detail="Не заполнены обязательные поля для полной формы")
        
        full_docx_path = main(
            docx_path=data["docx_path"],
            recipient=data["recipient"],
            signer=data["signer"],
            contract_number=data.get("contract_number"),
            contract_date=data.get("contract_date"),
            pdf_folder_path=data.get("pdf_folder_path")
        )
        print(f"[2/6] Обработанный путь: {full_docx_path}")
        
        
        return JSONResponse({
            "full_docx_path": full_docx_path,
            "status": "success"
        })
        
    except HTTPException:
        raise
    except Exception as e:
        error_msg = f"Непредвиденная ошибка: {str(e)}"
        print(f"[CRITICAL ERROR] {error_msg}")
        raise HTTPException(status_code=500, detail=error_msg)
    
@app.post("/api/extract-fields")
async def extract_fields(request: Request):
    """Извлекает поля вида {field_name} из DOCX."""
    try:
        data = await request.json()
        docx_path = data.get("docx_path")
        
        if not docx_path:
            raise HTTPException(status_code=400, detail="Не указан путь к документу")
        
        doc = Document(docx_path)
        fields = set()
        
        
        for paragraph in doc.paragraphs:
            fields.update(re.findall(r"\{(\w+)\}", paragraph.text))
    
        for section in doc.sections:
            for paragraph in section.header.paragraphs:
                fields.update(re.findall(r"\{(\w+)\}", paragraph.text))
            for paragraph in section.footer.paragraphs:
                fields.update(re.findall(r"\{(\w+)\}", paragraph.text))
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    matches = re.findall(r"\{(\w+)\}", cell.text)
                    fields.update(matches)

        print(fields)
        
        return {"fields": list(fields)}
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка при извлечении полей: {str(e)}")
    
@app.post("/api/replace-fields")
async def replace_fields(request: Request):
    try:
        data = await request.json()
        docx_path = data.get("docx_path")
        fields_data = data.get("fields_data", {})
        
        if not docx_path:
            raise HTTPException(status_code=400, detail="Не указан путь к документу")
        
        if not os.path.exists(docx_path):
            raise HTTPException(status_code=404, detail="Файл не найден")
        
        doc = Document(docx_path)
        

        for paragraph in doc.paragraphs:
            for field, value in fields_data.items():
                if f"{{{field}}}" in paragraph.text:
                    paragraph.text = paragraph.text.replace(f"{{{field}}}", str(value))
        
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for field, value in fields_data.items():
                        if f"{{{field}}}" in cell.text:
                            cell.text = cell.text.replace(f"{{{field}}}", str(value))
                            for paragraph in cell.paragraphs:
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                for run in paragraph.runs:
                                    run.font.italic = True
        
        
        
        doc.save(docx_path)
        
        return {"status": "success", "message": "Файл успешно обновлен", "file_path": docx_path}
    
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка при замене полей: {str(e)}")

@app.post("/api/add-table-row")
async def add_table_row(request: Request):
    data = await request.json()
    doc = Document(data["docx_path"])
    table = doc.tables[data["table_index"]]
    
    row_number = len(table.rows)
    new_row = table.add_row()
    
    if len(new_row.cells) >= 4:
        new_row.cells[0].text = str(row_number)
        new_row.cells[1].text = "{engineer_full_name}"
        new_row.cells[2].text = "{engineer_notebook}"
        new_row.cells[3].text = "{areas_name}"
        
        for cell in new_row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(40)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.save(data["docx_path"])
    return {"status": "success", "file_path": data["docx_path"]}

@app.get("/api/engineers")
async def get_engineers():
    engineers = [
        {"full_name": "Иванов И.И.", "notebook": "Lenovo X1"},
        {"full_name": "Петров П.П.", "notebook": "Dell XPS"},
        {"full_name": "Сидоров С.С.", "notebook": "MacBook Pro"}
    ]
    return {"engineers": engineers}

@app.get("/api/areas")
async def get_areas():
    areas = ["Зона A", "Зона B", "Зона C", "Зона D"]
    return {"areas": areas}