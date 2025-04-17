import os
import fitz
import re
import pdfplumber
from docx import Document
from typing import Dict, List
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

class PDFDataExtractor:
    def __init__(self):
        self.parsed_data = {}
        self.all_documents = []

    def extract_text_from_pdf(self, pdf_path: str) -> tuple:
        with fitz.open(pdf_path) as doc:
            first_page_text = doc[0].get_text("text")
            last_page_text = doc[-1].get_text("text") if len(doc) > 1 else ""
            num_pages = len(doc)
        return first_page_text, last_page_text, num_pages

    def format_title(self, title):
        words = title.split()
        formatted_words = []
        inside_quotes = False  
        quote_buffer = []  

        for i, word in enumerate(words):
            if word in {"ООО", "АО", "ЗАО", "ПАО", "ОАО"}:
                formatted_words.append(word)  
            elif i == 0:
                formatted_words.append(word.capitalize())  
            elif "«" in word: 
                inside_quotes = True
                quote_buffer = [word[1:].lower()]  
            elif "»" in word:  
                quote_buffer.append(word[:-1].lower())  
                formatted_words.append(f'«{" ".join(quote_buffer)}»')
                inside_quotes = False
            elif inside_quotes:
                quote_buffer.append(word.lower())
            else:
                formatted_words.append(word.lower()) 
        
        return " ".join(formatted_words)

    def clean_text(self, text: str) -> str:       
        lines = text.split("\n")
        cleaned_lines = [line.strip() for line in lines if line.strip()]  
        return "\n".join(cleaned_lines) 

    def extract_title_and_text(self, pdf_path: str) -> tuple:
        with pdfplumber.open(pdf_path) as pdf:
            page = pdf.pages[0]  
            lines = page.extract_text().split("\n")[:-1]  

            highlighted_lines = []
            is_title = False
            idx_ooo = 0
            for i, line in enumerate(lines):
                idx_ooo -= 1
                if line.isupper():
                    is_title = True
                if line.startswith("ООО"):
                    idx_ooo += i
                    start = line.find('«') + 1
                    end = line.find('»')                    
                    if start != -1 and end != -1:
                        quoted_text = line[start:end]                       
                        quoted_text = quoted_text[0].upper() + quoted_text[1:].lower()
                        text = line[:start] + quoted_text + line[end:]
                    highlighted_lines.append(text)
                    continue
                if line.isupper() and any(char.isdigit() for char in line):
                    highlighted_lines.append(line)
                    continue
                if line.isupper():
                    highlighted_lines.append(line.lower())
                    continue
                if not line.isupper() and is_title:
                    highlighted_lines.append(line)
        
        highlighted_lines[0] = highlighted_lines[0].capitalize()
        highlighted_lines[idx_ooo + 1] = highlighted_lines[idx_ooo + 1].capitalize()
        
        return " ".join(highlighted_lines[:idx_ooo + 1]), "\n".join(highlighted_lines[idx_ooo + 1:])

    def extract_kt_data(self, first_page: str, last_page: str, pdf_path: str) -> dict:
        title, remaining_text = self.extract_title_and_text(pdf_path)
        
        match_uch = re.search(r'Уч\. №\s*(\S+)', last_page)
        match_sheets = re.search(r'Отпечатан 1 экз\. на (\d+) л\.', last_page)
        match_date = re.search(r'Дата:\s*(.+)', last_page)
        match_exemplars = re.search(r'экз\. №№\s*(\S+)', last_page)
        
        return {
            'названиедокумента': title,
            'текстдокумента': remaining_text,
            'учетныйномер': match_uch.group(1).strip() if match_uch else "",
            'количестволистов': match_sheets.group(1).strip() if match_sheets else "",
            'датадокумента': match_date.group(1).strip() if match_date else "",
            'экземпляры': match_exemplars.group(1).strip() if match_exemplars else "",
            'типдокумента': 'КТ'
        }

    def extract_sp_data(self, first_page: str, last_page: str, num_pages: int, pdf_path: str) -> dict:
        title, remaining_text = self.extract_title_and_text(pdf_path)
        
        match_inv = re.search(r'Инв\. № подл\.\s*(\S+) ', first_page)
        
        return {
            'названиедокумента': title,
            'текстдокумента': remaining_text,
            'инвентарныйномер': match_inv.group(1).strip() if match_inv else "",
            'количествостраниц': num_pages,
            'типдокумента': 'СП'
        }

    def parse_pdf(self, pdf_path: str) -> dict:
        first_page, last_page, num_pages = self.extract_text_from_pdf(pdf_path)
        is_commercial_secret = "Коммерческая тайна" in first_page
        
        if is_commercial_secret:
            data = self.extract_kt_data(first_page, last_page, pdf_path)
        else:
            data = self.extract_sp_data(first_page, last_page, num_pages, pdf_path)
        
        self.all_documents.append(data)
        return data

    def generate_documents_list(self):
        ks_list = []
        fixed_item = "Электронная версия рабочей документации по проекту «Дооснащение подсистем безопасности кошек в обычной жизни ООО «Мяу мышь» на электронном носителе CD-R, коммерческая тайна, уч. № КТ/Э 459 от 26.03.2025(календарь), экз. № 1/1(человек должен задать отдельно для дисков и бумажной документации) только в адрес. "
        ks_list.append(fixed_item)
        
        for doc in self.all_documents:
            if doc['типдокумента'] == 'КТ':
                ks_entry = f"{doc['названиедокумента']}, Том 1, коммерческая тайна, Уч. № {doc['учетныйномер']} от {doc['датадокумента']}, экз. №№ {doc['экземпляры']} на {doc['количестволистов']} л. каждый, только в адрес."
                ks_list.append(ks_entry)
            elif doc['типдокумента'] == 'СП':
                sp_entry = f"{doc['названиедокумента']}, инв. № {doc['инвентарныйномер']} - на {doc['количествостраниц']} л. в 4 экз."
                ks_list.append(sp_entry)

        return ks_list

class DocxTemplateProcessor:
    def __init__(self):
        self.data_sources = {
            'парсера': {},
            'бд': {},
            'списокдокументов': ""
        }

    def add_parser_data(self, data: dict):
        """Добавляет данные из парсера."""
        self.data_sources['парсера'] = data

    def add_db_data(self, data: dict):
        """Добавляет данные из базы данных."""
        self.data_sources['бд'] = data

    def add_documents_list(self, documents_list: str):
        """Добавляет список документов."""
        self.data_sources['списокдокументов'] = documents_list
        
    def process_template(self, template_path: str, output_path: str):
        """Обрабатывает шаблон с выравниванием по ширине"""
        print("qwetryrte")
        doc = Document(template_path)
        print("sdfsdfsd")
        LEFT_INDENT = 1065
        HANGING_INDENT = 357
        FIRST_LINE = -357
        LINE_SPACING = 360
        
        for paragraph in doc.paragraphs:
            if '{списокдокументов:}' in paragraph.text:
                parent = paragraph._p.getparent()
                index = parent.index(paragraph._p)
                parent.remove(paragraph._p)
                
                for item in self.data_sources['списокдокументов']:
                    new_paragraph = OxmlElement('w:p')
                    
                    num_pr = OxmlElement('w:numPr')
                    ilvl = OxmlElement('w:ilvl')
                    ilvl.set(qn('w:val'), '0')
                    num_id = OxmlElement('w:numId')
                    num_id.set(qn('w:val'), '1')
                    num_pr.append(ilvl)
                    num_pr.append(num_id)
                    
                    p_pr = OxmlElement('w:pPr')
                    p_pr.append(num_pr)
                    
                    jc = OxmlElement('w:jc')
                    jc.set(qn('w:val'), 'both')
                    p_pr.append(jc)
                    
                    ind = OxmlElement('w:ind')
                    ind.set(qn('w:left'), str(LEFT_INDENT))
                    ind.set(qn('w:hanging'), str(HANGING_INDENT))
                    ind.set(qn('w:firstLine'), str(FIRST_LINE))
                    p_pr.append(ind)
                    
                    spacing = OxmlElement('w:spacing')
                    spacing.set(qn('w:before'), '0')
                    spacing.set(qn('w:after'), '0')
                    spacing.set(qn('w:line'), str(LINE_SPACING))
                    spacing.set(qn('w:lineRule'), 'auto')
                    p_pr.append(spacing)
                    
                    new_paragraph.append(p_pr)
                    
                    run = OxmlElement('w:r')
                    text = OxmlElement('w:t')
                    text.text = item
                    run.append(text)
                    new_paragraph.append(run)
                    
                    parent.insert(index, new_paragraph)
                    index += 1
                    
                break
        
        for paragraph in doc.paragraphs:
            self._replace_in_paragraph(paragraph)
            
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph)
        print("sdfsdfsd")
        doc.save(output_path)


    def _process_headers_footers(self, doc):
        """Обрабатывает верхние и нижние колонтитулы всех секций."""
        for section in doc.sections:
            for header_part in [section.header, section.footer, 
                              section.even_page_header, section.even_page_footer,
                              section.first_page_header, section.first_page_footer]:
                if header_part is not None:
                    for paragraph in header_part.paragraphs:
                        self._replace_in_paragraph(paragraph)
                    for table in header_part.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                for paragraph in cell.paragraphs:
                                    self._replace_in_paragraph(paragraph)

    def _replace_in_paragraph(self, paragraph):
        """Заменяет ключи в параграфе с сохранением форматирования."""
        keys = self._find_all_keys(paragraph.text)
        if not keys:
            return
       
        runs = paragraph.runs
        if not runs:
            return

        full_text = ''.join(run.text for run in runs)
        if not full_text:
            return

       
        for key, value in keys.items():
            source = value['source']
            key_name = value['key']
            if source == key_name and source in self.data_sources:
                replacement = str(self.data_sources[source])
            # Для сложных ключей
            elif source in self.data_sources and isinstance(self.data_sources[source], dict) \
                and key_name in self.data_sources[source]:
                replacement = str(self.data_sources[source][key_name])
            else:
                continue
            
        full_text = full_text.replace(key, replacement)
          

        for run in runs:
            run.text = ""

        if runs:
            runs[0].text = full_text
            for run in runs[1:]:
                self._copy_run_formatting(runs[0], run)

    def _copy_run_formatting(self, source_run, target_run):
        """Копирует форматирование из одного run в другой."""
        target_run.bold = source_run.bold
        target_run.italic = source_run.italic
        target_run.underline = source_run.underline
        target_run.font.name = source_run.font.name
        target_run.font.size = source_run.font.size
        if source_run.font.color.rgb:
            target_run.font.color.rgb = source_run.font.color.rgb

    def _find_all_keys(self, text: str) -> Dict[str, dict]:        
        keys = {}

        pattern = re.compile(r'\{([^}_]+)_из_([^}]+)\}')
        pattern_simple = re.compile(r'\{([^}]+)\:\}')
        
        for match in pattern.finditer(text):
            full_key = match.group(0)
            key_name = match.group(1)
            source = match.group(2)

            keys[full_key] = {
                'found': True,
                'source': source,
                'key': key_name
            }

        for match in pattern_simple.finditer(text):            
            full_key = match.group(0)
            key_name = match.group(1)
                     
            for source_name, source_data in self.data_sources.items():
                if isinstance(source_data, List) and source_name == key_name:
                    keys[full_key] = {
                        'found': True,
                        'source': source_name,
                        'key': key_name
                    }
                    break
       
        return keys

def main(docx_path):
    pdf_extractor = PDFDataExtractor()
    docx_processor = DocxTemplateProcessor()

    pdf_files = []
    for root, _, files in os.walk(r"C:\Users\andre\Desktop\Test"):
        pdf_files.extend([os.path.join(root, f) for f in files if f.lower().endswith('.pdf')])
    
    
    selected_indices = list(range(9))
    
    for idx in selected_indices:
        selected_pdf = pdf_files[idx]
        print(f"\nПарсинг файла: {selected_pdf}")
        parsed_data = pdf_extractor.parse_pdf(selected_pdf)
        docx_processor.add_parser_data(parsed_data)
    
    documents_list = pdf_extractor.generate_documents_list()
    docx_processor.add_documents_list(documents_list)
    

    db_data = {
        'ФИО': 'Иванов Иван Иванович',
        'должность': 'Менеджер'
    }
    docx_processor.add_db_data(db_data)

    template_path = f"C:/Users/andre/Desktop/VIKKA{docx_path}"
    print(template_path)
    output_path = r"converted_files\template.docx"
    docx_processor.process_template(template_path, output_path)
    print(f"\nДокумент сохранен как {output_path}")
    return output_path

if __name__ == "__main__":
    main()